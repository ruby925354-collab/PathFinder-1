from fastapi import APIRouter, UploadFile, File
from openpyxl import load_workbook
from database import get_db_connection
from fastapi import HTTPException
router = APIRouter()

@router.post("/api/knowledge-questions/import-excel")
async def import_knowledge_questions_excel(file: UploadFile = File(...)):
    # Save temporary file
    contents = await file.read() 
    temp_file = "uploaded_questions.xlsx"

    with open(temp_file, "wb") as f:
        f.write(contents)

    # Load workbook
    wb = load_workbook(temp_file, data_only=True)
    sheet = wb.active

    REQUIRED_COLUMNS = [
        "knowledge_type",
        "category",
        "category_type",
        "question",
        "option1",
        "option2",
        "option3",
        "option4",
        "correct_answer",
        "timer"
    ]

    # Read and normalize header row
    header = [
        str(cell.value).strip().lower() if cell.value else ""
        for cell in sheet[1]
    ]

    # Validate missing required columns
    missing_columns = [col for col in REQUIRED_COLUMNS if col.lower() not in header]

    if missing_columns:
        return {
            "status": "error",
            "message": f"Missing required header columns: {', '.join(missing_columns)}",
            "imported": 0,
            "errors": []
        }

    # Map header index positions to ensure proper column reading
    col_index = {col: header.index(col.lower()) for col in REQUIRED_COLUMNS}

    imported_count = 0
    errors = []
    row_number = 2  # start counting from the first data row

    with get_db_connection() as conn:
        cursor = conn.cursor()

        for row in sheet.iter_rows(min_row=2):
            try:
                # Read row values safely
                values = [cell.value for cell in row]

                # Extract and normalize row
                def get(col_name):
                    idx = col_index[col_name]
                    val = values[idx]
                    return str(val).strip() if val is not None else ""

                knowledge_type = get("knowledge_type")
                category = get("category")
                category_type = get("category_type")
                question = get("question")
                option1 = get("option1")
                option2 = get("option2")
                option3 = get("option3")
                option4 = get("option4")
                correct_answer = get("correct_answer")
                timer_value = get("timer")

                # REQUIRED FIELD VALIDATION
                required_values = [
                    ("knowledge_type", knowledge_type),
                    ("category", category),
                    ("category_type", category_type),
                    ("question", question),
                    ("option1", option1),
                    ("option2", option2),
                    ("option3", option3),
                    ("option4", option4),
                    ("correct_answer", correct_answer),
                ]

                missing_values = [name for name, val in required_values if val == ""]

                if missing_values:
                    raise ValueError(
                        f"Missing values in required columns: {', '.join(missing_values)}"
                    )

                # VALIDATE CORRECT ANSWER EXISTS IN OPTIONS
                options_list = [option1, option2, option3, option4]
                if correct_answer not in options_list:
                    raise ValueError(
                        f"Correct answer '{correct_answer}' is not one of the options"
                    )

                # TIMER VALIDATION
                try:
                    timer = int(float(timer_value)) if timer_value else 60
                except:
                    raise ValueError("Timer must be a number")

                # INSERT INTO knowledge_test
                cursor.execute("""
                    INSERT INTO knowledge_test 
                    (knowledge_type, category, question, answer, timer, category_type)
                    VALUES (%s, %s, %s, %s, %s, %s)
                """, (
                    knowledge_type,
                    category,
                    question,
                    correct_answer,
                    timer,
                    category_type
                ))

                knowledge_id = cursor.lastrowid

                # Insert each option as a separate row
                for opt in options_list:
                    cursor.execute("""
                        INSERT INTO knowledge_options (knowledge_id, choises)
                        VALUES (%s, %s)
                    """, (knowledge_id, opt))


                imported_count += 1

            except Exception as e:
                errors.append({
                    "row": row_number,
                    "error": str(e)
                })

            row_number += 1

        # Stop import if more than 0 fatal errors
        if errors:
            conn.rollback()
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "Import stopped due to formatting errors.",
                    "imported": imported_count,
                    "errors": errors
                }
            )

        conn.commit()

    return {
        "status": "success",
        "message": f"Successfully imported {imported_count} questions.",
        "imported": imported_count,
        "errors": errors
    }
